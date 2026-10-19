"""通用 Word 导出执行图：produce(LLM) -> make_docx（Markdown 轻解析排版）。

供纯文本产出的工具（如旅行计划）复用：工具在后台把 executor 设为 `docx_export`
即可获得 .docx 文件下载能力，无需各自写图。
"""
import io
import re
from typing import TypedDict

from langgraph.graph import StateGraph, END

from app.core.llm.gateway import LLMGateway


class ExportState(TypedDict, total=False):
    input: str
    prompt: str
    result: str
    model_used: str
    result_ref: str
    gateway: object
    inv_id: int


def _produce(s: ExportState) -> ExportState:
    gw: LLMGateway = s["gateway"]
    resp = gw.call(
        "executor",
        [
            {"role": "system", "content": s["prompt"]},
            {"role": "user", "content": s["input"]},
        ],
        max_tokens=4000,
    )
    s["result"] = resp["content"]
    s["model_used"] = resp["model_used"]
    return s


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _add_runs(par, text: str) -> None:
    """处理 **加粗** 标记，其余原样。"""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        run = par.add_run(m.group(1))
        run.bold = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def _is_table_line(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def _add_table(doc, lines: list) -> None:
    """把连续的 Markdown 表格行渲染为 docx 表格（跳过分隔行）。"""
    from docx.shared import Pt

    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue  # 分隔行 |---|---|
        rows.append(cells)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, r in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = r[j] if j < len(r) else ""
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
    doc.add_paragraph()


def _make_docx(s: ExportState) -> ExportState:
    """把 LLM 产出的 Markdown 轻解析后导出 docx（标题/列表/表格/加粗），上传 S3。"""
    from docx import Document
    from docx.shared import Pt

    from app.services import attachment as att

    doc = Document()
    # AIGC 隐式标识：文件元数据（§16.3-A，《AI 生成合成内容标识办法》）
    cp = doc.core_properties
    cp.author = "万宝工具箱 AI"
    cp.comments = "本文件由 AI（万宝工具箱小程序）生成，内容仅供参考"
    cp.category = "AI生成内容"

    lines = (s.get("result") or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if _is_table_line(stripped):
            block = []
            while i < len(lines) and _is_table_line(lines[i].strip()):
                block.append(lines[i])
                i += 1
            _add_table(doc, block)
            continue
        if not stripped:
            i += 1
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
        elif re.match(r"^[-*•]\s+", stripped):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", stripped))
        elif re.match(r"^\d+[.、)]\s+", stripped):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.、)]\s+", "", stripped))
        else:
            _add_runs(doc.add_paragraph(), stripped)
        i += 1

    # AIGC 显式标识：文末说明行（§16.3-A）
    note = doc.add_paragraph()
    run = note.add_run("—— 本文件由 AI（万宝工具箱小程序）生成，内容仅供参考 ——")
    run.font.size = Pt(9)
    run.italic = True
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    key = f"results/{s.get('inv_id')}.docx"
    att.upload(key, buf.getvalue())
    # 存 object_key，下载 URL 由 API 层代理回源（bucket 私有）
    s["result_ref"] = key
    return s


def build(tool) -> "CompiledGraph":  # noqa: F821
    g = StateGraph(ExportState)
    g.add_node("produce", _produce)
    g.add_node("make_docx", _make_docx)
    g.set_entry_point("produce")
    g.add_edge("produce", "make_docx")
    g.add_edge("make_docx", END)
    return g.compile()

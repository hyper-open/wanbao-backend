"""附件存储与解析（S3 兼容 rustfs 复用；docx/pdf/txt 文本抽取）。

M3：用户上传简历等附件 -> 存 S3 -> 解析出文本 -> 喂给工具执行。
解析为厚活，统一在异步队列 worker 中完成；此处提供解析原语。
"""
import io
import uuid
from typing import Optional

from sqlalchemy.orm import Session

from app.core.config import settings
from app.models import Attachment


def _client():
    import boto3

    return boto3.client(
        "s3",
        endpoint_url=settings.S3_ENDPOINT,
        aws_access_key_id=settings.S3_ACCESS_KEY,
        aws_secret_access_key=settings.S3_SECRET_KEY,
        region_name=settings.S3_REGION,
    )


def ensure_bucket():
    c = _client()
    try:
        c.head_bucket(Bucket=settings.S3_BUCKET)
    except Exception:  # noqa: BLE001
        c.create_bucket(Bucket=settings.S3_BUCKET)


def upload(object_key: str, data: bytes) -> str:
    _client().put_object(Bucket=settings.S3_BUCKET, Key=object_key, Body=data)
    return object_key


def get_object_bytes(object_key: str) -> bytes:
    resp = _client().get_object(Bucket=settings.S3_BUCKET, Key=object_key)
    return resp["Body"].read()


def get_url(object_key: str) -> str:
    return f"{settings.S3_ENDPOINT}/{settings.S3_BUCKET}/{object_key}"


def is_image(att: Attachment) -> bool:
    """附件是否为图片（多模态输入；与上传白名单里的 image/* 对齐）。"""
    mime = (att.mime or "").lower()
    low = (att.filename or "").lower()
    return mime.startswith("image/") or low.endswith((".jpg", ".jpeg", ".png", ".webp"))


def image_data_url(att: Attachment) -> Optional[str]:
    """图片附件 -> base64 data URL（视觉模型消费）。

    用 data URL 而非 presigned URL：S3 桶私有且公网可达性取决于 S3_PUBLIC_ENDPOINT
    配置，base64 对所有 OpenAI 协议供应商都可用；超限返回 None（调用方跳过并提示）。
    """
    import base64

    data = get_object_bytes(att.object_key)
    if len(data) > settings.ATTACHMENT_IMAGE_MAX_BYTES:
        return None
    mime = att.mime or "image/jpeg"
    if not mime.startswith("image/"):
        mime = "image/jpeg"
    return f"data:{mime};base64,{base64.b64encode(data).decode()}"


def image_parts(db: Session, object_keys: list) -> tuple[list[dict], list[str]]:
    """object_key 列表 -> (OpenAI 视觉 image_url parts, 被跳过的文件名)。

    空列表返回 ([], [])；全部超限时 parts 为空，调用方回落纯文本并附提示。
    """
    parts, skipped = [], []
    for key in object_keys or []:
        att = db.query(Attachment).filter(Attachment.object_key == key).first()
        if not att:
            continue
        url = image_data_url(att)
        if not url:
            skipped.append(att.filename or key)
            continue
        parts.append({"type": "image_url", "image_url": {"url": url}})
    return parts, skipped


def with_images(text: str, parts: list[dict], skipped: list[str]) -> "str | list[dict]":
    """组装多模态 user content：无可用图片时回落纯文本（附跳过说明）。"""
    if not parts:
        if skipped:
            note = "、".join(skipped)
            return text + "\n\n（附件 " + note + " 超过图片大小上限，未能识别）"
        return text
    return [{"type": "text", "text": text}] + parts


def presign_get_url(object_key: str, expires: int | None = None) -> str:
    """签发带签名的临时下载 URL（bucket 私有，裸 URL 会 403）。

    注意：SigV4 签名包含 Host，若 S3_PUBLIC_ENDPOINT 与内部 S3_ENDPOINT 不同，
    必须用对外地址签（签名才能匹配外网 Host）；该地址需反向代理到同一对象存储。
    """
    import boto3

    endpoint = settings.S3_PUBLIC_ENDPOINT or settings.S3_ENDPOINT
    c = boto3.client(
        "s3",
        endpoint_url=endpoint,
        aws_access_key_id=settings.S3_ACCESS_KEY,
        aws_secret_access_key=settings.S3_SECRET_KEY,
        region_name=settings.S3_REGION,
    )
    return c.generate_presigned_url(
        "get_object",
        Params={"Bucket": settings.S3_BUCKET, "Key": object_key},
        ExpiresIn=expires or settings.PRESIGN_EXPIRES_SECONDS,
    )


def save_upload(db: Session, user_id: int, filename: str, data: bytes, mime: str) -> Attachment:
    """落库 + 上传 S3，返回 Attachment 记录。"""
    key = f"attachments/{user_id}/{uuid.uuid4().hex}_{filename}"
    upload(key, data)
    att = Attachment(
        user_id=user_id,
        object_key=key,
        filename=filename,
        mime=mime,
        size=len(data),
        status="uploaded",
    )
    db.add(att)
    db.commit()
    db.refresh(att)
    return att


def _parse_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for tbl in doc.tables:
        for row in tbl.rows:
            parts.append(" | ".join(c.text for c in row.cells if c.text.strip()))
    return "\n".join(parts)


def _parse_pdf(data: bytes) -> str:
    import pdfplumber

    out = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            if t.strip():
                out.append(t)
    return "\n".join(out)


def _parse_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = []
    for si, slide in enumerate(prs.slides):
        parts = [f"[Slide {si + 1}]"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = "\n".join(
                    p.text for p in shape.text_frame.paragraphs if p.text.strip()
                )
                if txt.strip():
                    parts.append(txt)
        if len(parts) > 1:
            out.append("\n".join(parts))
    return "\n\n".join(out)


def parse_text(object_key: str, mime: str, filename: str = "") -> str:
    """抽取附件纯文本（按 mime / 扩展名路由）。"""
    data = get_object_bytes(object_key)
    low = (filename or "").lower()
    if mime == "application/pdf" or low.endswith(".pdf"):
        return _parse_pdf(data)
    if (
        mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        or low.endswith(".pptx")
    ):
        return _parse_pptx(data)
    if (
        mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or low.endswith(".docx")
    ):
        return _parse_docx(data)
    if mime in ("text/plain", "text/markdown") or low.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore")
    if mime.startswith("image/") or low.endswith((".jpg", ".jpeg", ".png", ".webp")):
        # 图片附件无文本可抽取：空串占位，待多模态模型直接消费图片
        return ""
    # 兜底：按文本尝试
    try:
        return data.decode("utf-8", errors="ignore")
    except Exception:  # noqa: BLE001
        return ""


def clip_text(text: str, limit: int) -> str:
    """超长文本截断并标注。

    附件解析文本会整段进 prompt（每轮都发），不设上限时一份大简历就能把
    单次请求撑爆；这里统一在「组装 prompt 前」截断，原文仍保留在 parsed_meta。
    """
    if not text:
        return ""
    if limit <= 0 or len(text) <= limit:
        return text
    return text[:limit] + f"\n…[内容过长已截断，仅保留前 {limit} 字]"


def ensure_parsed(db: Session, att: Attachment) -> str:
    """确保附件已解析：未解析则解析并写回 parsed_meta。返回文本。"""
    if att.parsed_meta and att.parsed_meta.get("text"):
        return att.parsed_meta["text"]
    text = parse_text(att.object_key, att.mime, att.filename or "")
    att.parsed_meta = {"text": text, "chars": len(text)}
    att.status = "parsed"
    db.commit()
    return text

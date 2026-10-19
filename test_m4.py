"""M4 端到端验证脚本（p313 + 复用中间件）。"""
import io
import json
import time
import urllib.request

import pptx
from pptx.util import Inches, Pt

BASE = "http://127.0.0.1:28082"


def _admin_token():
    """后台鉴权：账号登录换 Bearer 令牌（旧共享令牌已移除）。"""
    body = json.dumps({"username": "admin", "password": __import__("os").environ.get("ADMIN_PASSWORD", "change-me-in-prod")}).encode()
    r = urllib.request.Request(
        BASE + "/api/admin/login",
        data=body,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(r, timeout=30) as resp:
        return json.loads(resp.read().decode())["token"]


def _req(method, path, data=None, token=None, raw=None, ctype=None, headers=None):
    url = BASE + path
    h = dict(headers or {})
    if token:
        h["Authorization"] = "Bearer " + token
    if raw is not None:
        body = raw
        if ctype:
            h["Content-Type"] = ctype
    elif data is not None:
        body = json.dumps(data).encode()
        h["Content-Type"] = "application/json"
    else:
        body = None
    req = urllib.request.Request(url, data=body, method=method, headers=h)
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return r.status, r.read().decode()
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode()


def make_pptx(path):
    prs = pptx.Presentation()
    blank = prs.slide_layouts[6]
    # slide 1 title
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1.2))
    tb.text_frame.text = "我的季度汇报"
    # slide 2 content
    s2 = prs.slides.add_slide(blank)
    tb2 = s2.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1.0))
    tb2.text_frame.text = "工作亮点"
    cb = s2.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(9), Inches(4))
    ctf = cb.text_frame
    ctf.text = "• 完成项目A\n• 推进项目B\n• 优化流程"
    prs.save(path)


def main():
    print("== 1. 登录 ==")
    st, body = _req("POST", "/api/auth/login", {"code": "test_code"})
    print(st, body)
    tok = json.loads(body)["token"]
    openid = json.loads(body).get("openid")
    h = {"Authorization": f"Bearer {tok}"}

    print("\n== 1b. 充值获取次数 ==")
    st, body = _req("GET", "/api/billing/packages", headers=h)
    pkg_id = json.loads(body)[0]["id"]
    st, body = _req("POST", "/api/billing/recharge/create", {"package_id": pkg_id}, headers=h)
    order_id = json.loads(body)["order_id"]
    st, body = _req("POST", "/api/payment/notify", {"order_id": order_id})
    print("recharge:", st, body)
    st, body = _req("GET", "/api/billing/account/me", headers=h)
    print("account:", body)

    print("\n== 2. 内容安全：正常识别 ==")
    st, body = _req("POST", "/api/agent/recognize", {"text": "帮我优化一下这份简历"}, headers=h)
    print(st, body[:160])

    print("\n== 3. 内容安全：命中拦截词 ==")
    st, body = _req("POST", "/api/agent/recognize", {"text": "这是一段包含违禁词示例A的测试"}, headers=h)
    print("期望 400 ->", st, body[:120])

    print("\n== 4. PPT 视觉重排：上传 pptx ==")
    import tempfile, os
    fd, p = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    make_pptx(p)
    with open(p, "rb") as f:
        pptx_bytes = f.read()
    # multipart upload
    boundary = "----wbtest"
    parts = []
    parts.append(f"--{boundary}".encode())
    parts.append(b'Content-Disposition: form-data; name="files"; filename="test.pptx"')
    parts.append(b"Content-Type: application/vnd.openxmlformats-officedocument.presentationml.presentation")
    parts.append(b"")
    parts.append(pptx_bytes)
    parts.append(f"--{boundary}--".encode())
    parts.append(b"")
    body_raw = b"\r\n".join(parts)
    st, body = _req("POST", "/api/attachment/upload", raw=body_raw,
                    ctype=f"multipart/form-data; boundary={boundary}", headers=h)
    print(st, body[:160])
    att_ids = json.loads(body)["attachment_ids"]
    print("attachment_ids:", att_ids)

    print("\n== 5. confirm pptx_optimize (异步) ==")
    st, body = _req("POST", "/api/agent/confirm",
                    {"tool_id": "pptx_optimize", "input": "请让每页更精炼", "attachment_ids": att_ids},
                    headers=h)
    print(st, body[:200])
    inv_id = json.loads(body)["invocation_id"]

    print("\n== 6. 轮询结果 ==")
    for i in range(20):
        time.sleep(2)
        st, body = _req("GET", f"/api/agent/invocation/{inv_id}", headers=h)
        j = json.loads(body)
        print(f"  poll {i}: status={j['status']}")
        if j["status"] in ("success", "failed"):
            print("  result_ref:", j.get("result_ref"))
            print("  result_text:", (j.get("result_text") or "")[:200])
            break

    print("\n== 7. 审计日志（内容安全命中应留痕）==")
    st, body = _req("GET", "/api/admin/monitor", token=_admin_token())
    print(st)
    m = json.loads(body)
    print("monthly:", json.dumps(m["monthly"], ensure_ascii=False))
    print("audit_hits count:", len(m["audit_hits"]))
    if m["audit_hits"]:
        print("  sample hit:", json.dumps(m["audit_hits"][0], ensure_ascii=False)[:200])

    print("\n== 8. 管理后台监控页可访问 ==")
    st, body = _req("GET", "/admin")
    print("admin status:", st, "html?", "<html" in body.lower())

    print("\nALL DONE")


if __name__ == "__main__":
    main()

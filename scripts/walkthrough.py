#!/usr/bin/env python3
"""阶段 0 内部走查：40 工具 + 6 边界任务，全链路走 /api/chat 真实请求链。"""
import io
import json
import os
import sys
import time
import concurrent.futures as cf

import requests

BASE = os.environ.get("WB_BASE", "http://localhost:28082/api/agent")
UPLOAD = os.environ.get("WB_UPLOAD", "http://localhost:28082/api/attachment/upload")
TOKEN = sys.argv[1]
H = {"Authorization": f"Bearer {TOKEN}"}

# (case_id, layer, expected_tool_id, input, special)
# special: "" 普通 | "upload_resume" 上传简历再跑 | "upload_ppt" 上传PPT再跑
#          "clarify_ok" 允许一次澄清后按预期工具确认 | "vague" 模糊意图只观察
#          "followup" 续问两轮 | "empty" 空输入 | "export" 期望导出链接
CASES = [
    ("A1", "A", "translate", "帮我把这句话翻成英文：明天上午十点开会，请带上门禁卡", ""),
    ("A2", "A", "weekly_report", "帮我写周报，这周做了：修了登录 bug、面试了 2 个人、上线了活动页", ""),
    ("A3", "A", "resume_optimize", "优化我的简历，突出项目亮点", "upload_resume"),
    ("A4", "A", "xhs_copy", "写一篇小红书文案，推广我的社区咖啡店，主打手冲", ""),
    ("A5", "A", "travel_plan", "做个杭州两日游攻略，预算 1500，要含行程表", "export"),
    ("A6", "A", "recipe_idea", "家里有鸡蛋、番茄、剩米饭，做什么吃？", ""),
    ("A7", "A", "title_gen", "给一篇讲时间管理的文章起 10 个标题", ""),
    ("A8", "A", "pet_train", "狗两岁，见到陌生人就狂叫，怎么训练？", ""),
    ("A9", "A", "smart_reply", "领导说'辛苦了，再帮我改版 PPT 吧'，怎么回显得情商高？", ""),
    ("A10", "A", "observability", "我们有个 Spring Boot 微服务要上线，怎么搭日志/指标/链路追踪？", ""),
    ("B1", "B", "unit_test", "给这个函数写单元测试：def add(a,b): return a+b if a>0 else -a-b", ""),
    ("B2", "B", "elevator_pitch", "帮我把'宠物上门喂养平台'写成 30 秒电梯演讲", ""),
    ("B3", "B", "ai_prompt", "我想要'赛博朋克风格的猫'的 Midjourney 提示词", ""),
    ("B4", "B", "folk_custom", "端午节除了吃粽子还有什么习俗？给孩子讲讲由来", ""),
    ("B5", "B", "architecture_review", "评审一下这个架构：前后端分离 + MySQL 主从 + Redis 缓存，日活 5 万", ""),
    ("B6", "B", "api_security", "我们的支付回调接口没做验签，帮我过一遍安全隐患", ""),
    ("B7", "B", "tech_debt", "老项目三年没大改，帮我做技术债务盘点计划", ""),
    ("B8", "B", "incident_postmortem", "昨天数据库慢查询导致下单超时 20 分钟，写个复盘报告", ""),
    ("B9", "B", "restaurant_plan", "想开家 30 平的米粉店，预算 20 万，给我筹备指南", ""),
    ("B10", "B", "zodiac_virgo", "处女座女朋友总嫌我乱放东西，怎么相处？", ""),
    ("C1", "C", "moving_checklist", "下周末搬家，帮我列个清单和时间表", ""),
    ("C2", "C", "logline", "把'一个机器人学会衰老'浓缩成一句话故事", ""),
    ("C3", "C", "interview_finjob", "下周面试财务分析岗，帮我准备常见面试题", ""),
    ("C4", "C", "fest_womens", "三八节快到了，给我个母婴店的营销活动方案", ""),
    ("C5", "C", "sop_write", "把门店收银流程写成 SOP", ""),
    ("C6", "C", "bug_fix", "排查这个报错：TypeError: cannot read properties of undefined (reading 'map')，React 列表页", ""),
    ("C7", "C", "color_scheme", "做一个健康类 App 的主配色，要清新", ""),
    ("C8", "C", "chart_advice", "我要展示 12 个月的销量对比和占比，用什么图？", ""),
    ("C9", "C", "ortho_teeth", "成人矫正牙齿大概要多少钱、有什么坑？", ""),
    ("C10", "C", "book_alive_b", "帮我深度解读一下《活着》这本书", ""),
    ("D1", "D", "pptx_optimize", "优化这份 PPT 的排版并导出", "upload_ppt"),
    ("D2", "D", "swot", "帮我用 SWOT 分析一下我的奶茶店加盟计划", "clarify_ok"),
    ("D3", "D", "ab_test_plan", "想测新首页 banner 是否提升点击，设计个 A/B 实验", "clarify_ok"),
    ("D4", "D", "mock_interview", "模拟面试我，产品经理岗，3 年经验", "clarify_ok"),
    ("D5", "D", "fitness_plan", "男，30 岁，久坐，想在家练出腹肌，给个计划", "clarify_ok"),
    ("D6", "D", "wechat_article", "写一篇'年轻人开始反向消费'的公众号文章", ""),
    ("D7", "D", "white_paper", "给我一份'AI 客服行业白皮书'的框架和核心章节内容", ""),
    ("D8", "D", "code_review", "审查这段代码的问题：\nfunction f(obj){ return obj.items.map(x=>x.name).join(','); }  // obj.items 可能为 undefined，且在大对象上循环引用", ""),
    ("D9", "D", "data_interpret", "解读这组数据：1月留存 40%，2月 35%，3月 30%，4月 28%，5月 27%，6月 26.5%", ""),
    ("D10", "D", "meeting_minutes", "把会议记录整理成纪要：今天讨论了下版本上线时间，老王说后端还要两天，小李说测试得三天，最后定下周五上线，另外老张要把文档补一下", ""),
    ("X1", "X", "recipe_idea", "帮我配一顿减脂餐", "vague"),
    ("X2", "X", None, "下周要见丈母娘，送什么好", "vague"),
    ("X3", "X", None, "帮我做个东西", "vague"),
    ("X4", "X", None, "写点东西", "vague"),
    ("X5", "X", "translate", "翻译：月色真美", "followup"),
    ("X6", "X", "travel_plan", "", "empty"),
]

DOC = b"""PK\x03\x04 placeholder"""  # noqa


def make_resume_docx() -> bytes:
    import docx as _d
    d = _d.Document()
    d.add_heading("张三", 0)
    d.add_paragraph("3 年后端开发，Python/FastAPI，做过电商订单系统与支付网关")
    d.add_paragraph("工作经历：某科技公司 2022-至今")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_ppt() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    for i in range(6):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"测试页 {i+1}"
        s.placeholders[1].text = "产品介绍与数据展示" * 3
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def api(method, path, **kw):
    kw.setdefault("timeout", 30)
    r = requests.request(method, BASE + path, headers=H, **kw)
    try:
        return r.status_code, r.json()
    except Exception:  # noqa: BLE001
        return r.status_code, {"raw": r.text[:300]}


def api_upload(files, data):
    r = requests.post(UPLOAD, headers=H, files=files, data=data, timeout=60)
    try:
        return r.status_code, r.json()
    except Exception:  # noqa: BLE001
        return r.status_code, {"raw": r.text[:300]}


def poll_session(sid, max_s=300):
    """轮询到终态；awaiting_confirm→confirm；awaiting_answer→answer 最多 2 次。

    注意：状态可能出现在会话顶层（status）或 active_turn 镜像，任一为准。
    """
    t0, answers = time.time(), 0
    last = {}
    while time.time() - t0 < max_s:
        code, data = api("GET", f"/session/{sid}")
        if code != 200:
            return {"error": f"poll {code}: {data}"}
        last = data
        t = data.get("active_turn") or {}
        cur = (t.get("status") or data.get("status") or "").lower()
        if cur == "awaiting_confirm":
            api("POST", f"/session/{sid}/confirm")
        elif cur == "awaiting_answer":
            if answers >= 2:
                api("POST", f"/session/{sid}/skip")
            else:
                answers += 1
                api("POST", f"/session/{sid}/answer", json={"text": "继续，按默认来"})
        elif cur in ("success", "completed", "failed", "cancelled", "awaiting_turn"):
            return data
        time.sleep(3)
    return {"error": "timeout", "last": {"status": last.get("status")}}


def run_case(case):
    cid, layer, expect, text, special = case
    out = {"case": cid, "layer": layer, "expect": expect, "hit": None,
           "router_status": None, "clarify_or_menu": None, "final": None,
           "result_len": 0, "export": None, "problem": None}
    try:
        if special == "upload_resume":
            _, up = api_upload(
                        {"files": ("resume.docx", make_resume_docx(),
                                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                        {"filename": "resume.docx"})
            if "attachment_ids" not in up:
                out["problem"] = f"P0 上传简历失败: {up}"
                return out
            code, res = api("POST", "/chat", json={"tool_id": "resume_optimize",
                                                   "input": text, "attachment_ids": up["attachment_ids"]})
        elif special == "upload_ppt":
            _, up = api_upload(
                        {"files": ("deck.pptx", make_ppt(),
                                          "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
                        {"filename": "deck.pptx"})
            if "attachment_ids" not in up:
                out["problem"] = f"P0 上传PPT失败: {up}"
                return out
            code, res = api("POST", "/chat", json={"tool_id": "pptx_optimize",
                                                   "input": text, "attachment_ids": up["attachment_ids"]})
        elif special == "empty":
            code, res = api("POST", "/chat", json={"tool_id": "travel_plan", "input": ""})
        else:
            code, res = api("POST", "/chat", json={"input": text})

        if code != 200:
            out["problem"] = f"HTTP {code}: {json.dumps(res, ensure_ascii=False)[:200]}"
            out["router_status"] = f"http{code}"
            return out
        out["router_status"] = res.get("status")

        # 路由层：menu/clarify/confirm/need_attachment 都不是已开会话
        if res.get("status") in ("clarify", "confirm"):
            out["hit"] = res.get("tool_id")
            out["clarify_or_menu"] = res.get("clarify") or "confirm"
            if special == "vague":
                return out  # 观察即可，不确认
            if res.get("tool_id") != expect:
                out["problem"] = f"路由偏差: 期望 {expect} 实际 {res.get('tool_id')}"
            code, res = api("POST", "/chat", json={"tool_id": res.get("tool_id"), "input": text})
            if code != 200 or res.get("status") != "session":
                out["problem"] = (out["problem"] or "") + f" 确认后开会话失败: {code} {json.dumps(res, ensure_ascii=False)[:150]}"
                return out
        elif res.get("status") == "menu":
            menu = [m.get("id") for m in (res.get("menu") or [])]
            out["clarify_or_menu"] = f"menu:{menu[:6]}"
            if special == "vague":
                return out
            out["hit"] = "IN_MENU" if expect in menu else "MISS"
            if expect not in menu:
                out["problem"] = f"菜单未包含期望工具 {expect}"
                return out
            code, res = api("POST", "/chat", json={"tool_id": expect, "input": text})
            if code != 200 or res.get("status") != "session":
                out["problem"] = (out["problem"] or "") + " 选定后开会话失败"
                return out
        elif res.get("status") == "need_attachment":
            out["hit"] = res.get("tool_id")
            out["clarify_or_menu"] = "need_attachment"
            return out
        elif res.get("status") != "session":
            out["hit"] = res.get("tool_id")
            out["problem"] = f"未知状态: {json.dumps(res, ensure_ascii=False)[:200]}"
            return out

        sid = res["session_id"]
        out["hit"] = res.get("tool_id") or out["hit"]
        data = poll_session(sid)
        if "error" in data:
            out["problem"] = f"轮询失败: {data['error']}"
            out["final"] = data.get("last", {}).get("status")
            return out
        t = data.get("active_turn") or {}
        if not t and data.get("turns"):
            t = data["turns"][-1]
        out["final"] = t.get("status") or data.get("status")
        if out["final"] == "awaiting_turn":
            out["final"] = "success_turn_end"
        if data.get("status") == "cancelled" and out["final"] not in ("success", "success_turn_end"):
            out["problem"] = f"会话被自动结束: {data.get('cancel_reason','')}"[:120]
        rt = t.get("result_text") or ""
        if not rt:
            rt = "".join((m.get("content") or "") for m in (t.get("messages") or [])
                         if m.get("role") == "assistant")
        out["result_len"] = len(rt)
        if out["final"] == "failed":
            err = t.get("error") or data.get("error") or ""
            out["problem"] = f"产出失败: {str(err)[:150]}"
        if special == "export" and rt:
            out["export"] = "docx链接" if ("docx" in rt.lower() or ".docx" in rt or "下载" in rt) else "无导出链接"
        return out
    except Exception as e:  # noqa: BLE001
        out["problem"] = f"异常: {type(e).__name__} {e}"
        return out


def run_followup():
    """X5 续问：同一会话先翻译，再续问。"""
    out = {"case": "X5", "layer": "X", "expect": "translate", "hit": None,
           "router_status": None, "clarify_or_menu": None, "final": None,
           "result_len": 0, "export": None, "problem": None}
    code, res = api("POST", "/chat", json={"input": "翻译：月色真美"})
    out["router_status"] = res.get("status")
    if res.get("status") in ("clarify", "confirm"):
        out["hit"] = res.get("tool_id")
        code, res = api("POST", "/chat", json={"tool_id": res.get("tool_id"), "input": "翻译：月色真美"})
    if code != 200 or res.get("status") != "session":
        out["problem"] = "首轮未开会话"
        return out
    poll_session(res["session_id"], max_s=180)
    # 续问：不带 tool_id，验证上一工具先验
    code, res2 = api("POST", f"/session/{res['session_id']}/turn", json={"input": "再来一句：今晚的月色也很美"})
    out["clarify_or_menu"] = res2.get("status")
    if code != 200:
        out["problem"] = f"续问 HTTP {code}: {json.dumps(res2, ensure_ascii=False)[:150]}"
        return out
    out["hit"] = res2.get("tool_id")
    if res2.get("status") == "session":
        data = poll_session(res2["session_id"], max_s=180)
        t = data.get("active_turn") or {}
        out["final"] = t.get("status")
        rt = t.get("result_text") or "".join((m.get("content") or "") for m in (t.get("messages") or []) if m.get("role") == "assistant")
        out["result_len"] = len(rt)
        if out["final"] == "failed":
            out["problem"] = "续问产出失败: " + str(t.get("error"))[:120]
    return out


def main():
    only = set(sys.argv[2].split(",")) if len(sys.argv) > 2 else None
    cases = [c for c in CASES if c[4] != "followup"]
    if only:
        cases = [c for c in cases if c[0] in only]
    results = []
    if only is None or "X5" in only:
        r = run_followup()
        results.append(r)
        print(f"[X5] hit={r['hit']} final={r['final']} len={r['result_len']} "
              f"prob={r['problem'] or '-'}", flush=True)
        time.sleep(75)
    for c in cases:
        r = run_case(c)
        results.append(r)
        print(f"[{r['case']}] hit={r['hit']} final={r['final']} len={r['result_len']} "
              f"prob={r['problem'] or '-'}", flush=True)
        if len(cases) > 1:
            time.sleep(75)
    with open("/tmp/walkthrough_result.json", "w") as f:
        json.dump(results, f, ensure_ascii=False, indent=1)
    print("SAVED /tmp/walkthrough_result.json")


if __name__ == "__main__":
    main()
